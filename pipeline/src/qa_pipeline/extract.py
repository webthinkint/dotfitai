"""Extract text lines from .docx files (paragraph-level, hyperlinks included).

python-docx's ``paragraph.text`` omits text inside hyperlink runs, so we go one
level lower and join all ``w:t`` nodes of each paragraph. Tables are flattened
to pipe-joined rows so their content survives into Stage 2 as plain text.

Extraction never raises for corrupted files: the caller gets an ExtractResult
with ``ok=False`` so the file can be routed to the manual queue.
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class ExtractResult:
    ok: bool
    lines: list[str] = field(default_factory=list)
    error: str | None = None
    n_paragraphs: int = 0
    n_tables: int = 0


def _para_text(par: Paragraph) -> str:
    """Paragraph text including hyperlink runs, with <w:br/>/<w:cr/> as \n.

    Saved-email threads in this corpus embed the whole From/Sent/To/Subject
    header (and the Name/Email/Question contact block) in a single paragraph
    separated by soft line breaks — collapsing them onto one line destroys
    Stage 1's markers.
    """
    parts: list[str] = []
    for node in par._p.iter():
        if node.tag == f"{W_NS}t":
            parts.append(node.text or "")
        elif node.tag in (f"{W_NS}br", f"{W_NS}cr"):
            parts.append("\n")
        elif node.tag == f"{W_NS}tab":
            parts.append(" ")
    return "".join(parts)


def _table_rows(table: Table) -> list[str]:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cells.append(" ".join(_para_text(p) for p in cell.paragraphs).strip())
        rows.append("| " + " | ".join(cells) + " |")
    return rows


def _walk_body(doc: DocumentObject) -> list[str]:
    """Yield lines in document order, interleaving paragraphs and tables.

    Paragraph text is split on embedded line breaks so that each physical
    line becomes its own element (markers must live on their own lines).
    """
    from docx.oxml.ns import qn

    lines: list[str] = []
    tables = list(doc.tables)  # document order; nth w:tbl ↔ nth Table
    tbl_idx = 0
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            for piece in _para_text(Paragraph(child, doc)).split("\n"):
                lines.append(piece)
        elif child.tag == qn("w:tbl"):
            if tbl_idx < len(tables):
                lines.extend(_table_rows(tables[tbl_idx]))
            tbl_idx += 1
    return lines


def extract_docx(path: Path) -> ExtractResult:
    try:
        doc = Document(path)
        lines = _walk_body(doc)
        return ExtractResult(
            ok=True,
            lines=lines,
            n_paragraphs=len(doc.paragraphs),
            n_tables=len(doc.tables),
        )
    except (zipfile.BadZipFile, ValueError, KeyError, OSError) as exc:
        return ExtractResult(ok=False, error=f"{type(exc).__name__}: {exc}")
    except Exception as exc:  # unexpected — still route to manual queue, never crash the run
        return ExtractResult(ok=False, error=f"UNEXPECTED {type(exc).__name__}: {exc}")
