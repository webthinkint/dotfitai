"""dotFIT Product Summaries deck -> markdown sections (data/Product Summaries).

Owner-ruled 2026-09-28: the deck is legal-approved and the highest authority
for now; on a disagreement it wins, then the PDSRG. Prices are individualized
per customer on the website, so no price leaves this script — every currency
amount is masked to ``[price]`` before anything is written.

Extraction is text-exact (python-pptx), never model-written:

- shapes are read in position order (top, then left; a group's children in
  the group's own position order) — the deck's shape z-order is authoring
  order and interleaves captions on the collage slides;
- tables become markdown tables, hyperlinks are kept per slide;
- each slide records the md5 of every picture it places.

The deck has no PowerPoint sections and almost no title placeholders, so the
split into sections is the hand-curated ``SECTIONS`` map below (like the
PDSRG's ``STEM_META``): every slide is either in exactly one section or in
``DROPPED`` with a reason, and the run fails otherwise. Product names resolve
through the alias table and the run fails on one that does not — an
unresolvable tag must never produce untagged sections silently.

On the bundle slides the membership of each level is carried by bottle
photos, not text (there is no alt text anywhere in the deck). ``PICTURED``
holds a hand-written note of what each such slide shows, read off the
rendered slides on 2026-09-28; it is rendered under the slide's text,
labelled as an annotation so it is never mistaken for deck wording.

``kind`` per section: ``product`` (a product's summary), ``bundle`` (how
products are combined), ``brand`` (the dotFIT difference, credibility),
``education`` (nutrition teaching), ``script`` (trainer scripts, taglines,
flyer and email copy — owner-ruled 2026-09-28 to hold back for now, possibly
for a separate tool later).

Output: processed/summaries/md/NN-<slug>.md + summary.json (sorted, no
timestamps — reruns are byte-identical, the pipeline determinism rule), and
the full per-slide dump under processed/summaries/runs/ (gitignored).

Usage (from pipeline/): uv run scripts/summaries_pptx_md.py
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

DECK_NAME = "summaries_teaching.pptx"
PRICE = "[price]"

# Currency amounts and per-unit prices. Order matters: the unit forms first,
# so "$1.74/svg" masks once, not as "[price]/svg" plus a stray match.
_PRICE_PATTERNS = [
    re.compile(r"\$\s?\d[\d,]*(?:\.\d+)?\+?(?:\s?/\s?(?:svg|serving|meal|g|gm))?", re.I),
    re.compile(r"~?(?<![\w.])\d*\.?\d+\s?(?:¢|c(?=\s?/))"),       # 40¢, 2.99c/d, .39c/
    re.compile(r"\b\d+\.\d{2}\s?/\s?svg\b", re.I),                  # WS 1.39/svg
]

KINDS = ("product", "bundle", "brand", "education", "script")

# (slug, title, kind, slides, products). Products are alias-table family
# names. Deck order; a section may be non-contiguous (slide 15, a family
# health-cabinet picture, belongs with the family bundle at 256+).
SECTIONS: list[tuple[str, str, str, list[int], list[str]]] = [
    ("practitioner-difference", "The practitioner difference", "brand",
     [2, 3], ["AminoFormula"]),
    ("performance-products", "Performance products and how to bundle them", "bundle",
     [4, 5, 6], ["First String", "Creatine Monohydrate", "NO7 PreWorkout", "Active MV",
                 "Alln1 SuperBlend", "Omega-3 Fish Oil", "Creatine Complex"]),
    ("senior-playspan-bundles", "Senior health, performance & longevity bundles (Playspan)", "bundle",
     [7, 8, 9, 10, 11, 12, 13], ["Over 50 MV", "Alln1 SuperBlend", "WheySmooth", "Omega-3 Fish Oil",
                                 "Calcium Complex", "Vitamin D-3", "Creatine Monohydrate", "AminoFormula"]),
    ("dotbar-new-bars", "New high-fiber protein bars", "product",
     [14], ["dotBAR"]),
    ("extras-and-taglines", "Performance, recovery and health extras, with taglines", "script",
     [16, 17, 18], ["Creatine Complex", "NO7 PreWorkout", "GlutamineComplex", "Digestive Enzymes",
                    "SleepAid", "Electrolytes", "Alln1 SuperBlend", "Antioxidant", "Active MV",
                    "Probiotics", "CollagenComplex", "Omega-3 Fish Oil", "Calcium Complex", "Vitamin D-3"]),
    ("dotfit-difference", "Dietary supplements — the dotFIT difference", "brand",
     [19, 23, 24, 26, 27, 28, 31, 32], []),
    ("trainer-client-handout", "Nutrition & supplementation — trainer and client handout", "education",
     [1, 36, 37], ["Active MV", "Omega-3 Fish Oil", "Calcium Complex"]),
    ("multivitamins", "The dotFIT multivitamin & mineral formulas", "product",
     [39], ["Active MV", "Women's MV", "Over 50 MV", "Alln1 SuperBlend"]),
    ("activemv", "ActiveMV", "product",
     [40, 42, 43, 44, 45, 46, 49, 50], ["Active MV"]),
    ("activemv-scripts", "ActiveMV — trainer script and taglines", "script",
     [47, 48], ["Active MV"]),
    ("superblend", "Alln1 SuperBlend", "product",
     [51, 52, 53, 54, 55, 56], ["Alln1 SuperBlend", "WheySmooth", "Plant Protein", "Omega-3 Fish Oil",
                                "Calcium Complex"]),
    ("over50mv", "Over50MV", "product",
     [58, 59, 60], ["Over 50 MV", "Active MV"]),
    ("womensmv", "Women's MV", "product",
     [61, 62, 63], ["Women's MV", "Active MV"]),
    ("multivitamin-guide", "Multivitamin guide — which formula for whom", "product",
     [66], ["Active MV", "Women's MV", "Over 50 MV", "Alln1 SuperBlend"]),
    ("vitamin-d3", "Vitamin D3", "product",
     [67, 68, 69, 70, 72, 74], ["Vitamin D-3"]),
    ("vitamin-d3-scripts", "Vitamin D3 — trainer scripts", "script",
     [73], ["Vitamin D-3"]),
    ("calcium-complex", "Calcium Complex", "product",
     [75, 76, 77], ["Calcium Complex"]),
    ("omega-3", "SuperOmega-3 fish oil", "product",
     [78, 79, 80, 81, 82, 83, 84], ["Omega-3 Fish Oil"]),
    ("playspan-essentials-bundle", "Playspan essentials bundle", "bundle",
     [85, 86, 87], ["Active MV", "Women's MV", "Over 50 MV", "Alln1 SuperBlend", "Omega-3 Fish Oil",
                    "Calcium Complex", "WheySmooth", "First String", "Pre & Post Workout Formula",
                    "LeanMeal Nutrition Shake", "Plant Protein"]),
    ("road-to-health-part-1-2", "The Road to Health — Parts 1 & 2 summary", "education",
     [88, 89, 90, 91], []),
    ("road-to-health-email", "The Road to Health — email campaign version", "script",
     [92, 93, 94], []),
    ("collagen-complex", "Collagen Complex", "product",
     [95, 96, 97, 98, 99], ["CollagenComplex"]),
    ("antioxidant", "Antioxidants (SuperiorAntioxidant)", "product",
     [100, 101, 102, 103], ["Antioxidant"]),
    ("antioxidant-scripts", "Antioxidants — trainer script", "script",
     [104], ["Antioxidant"]),
    ("probiotics", "Probiotics", "product",
     [105, 106, 107, 108, 109], ["Probiotics"]),
    ("digestive-enzymes", "Digestive Enzymes", "product",
     [110, 111, 112, 113, 114, 115], ["Digestive Enzymes"]),
    ("sleep-aid", "Sleep Aid", "product",
     [116, 117, 118], ["SleepAid"]),
    ("brain-health", "Brain Health", "product",
     [119, 121, 122], ["Brain Health"]),
    ("brain-mental-health-levels", "Brain & mental health dietary support — levels 1 and 2", "bundle",
     [123, 124, 125], ["Active MV", "Alln1 SuperBlend", "Vitamin D-3", "Omega-3 Fish Oil",
                       "Creatine Monohydrate", "Brain Health"]),
    ("youth-preserving-bundles", "Youth-preserving nutrition bundles, levels 1–5", "bundle",
     [120, 126, 127, 128, 129], ["Active MV", "Women's MV", "Over 50 MV", "Alln1 SuperBlend",
                                 "Calcium Complex", "Omega-3 Fish Oil", "Vitamin D-3", "Antioxidant",
                                 "Probiotics", "CollagenComplex", "Brain Health", "Digestive Enzymes",
                                 "SleepAid", "LeanMeal Nutrition Shake", "WheySmooth"]),
    ("proteins", "dotFIT proteins — overview and lineup", "product",
     [130, 131, 132, 134, 161], ["WheySmooth", "Plant Protein", "First String", "LeanMeal Nutrition Shake",
                                 "Pre & Post Workout Formula"]),
    ("proteins-scripts", "dotFIT proteins — video/trainer script", "script",
     [133], ["WheySmooth", "First String"]),
    ("plant-protein", "Plant Protein (BestPlantProtein)", "product",
     [135, 136, 137], ["Plant Protein"]),
    ("pre-post-workout", "Pre & Post Workout Formula", "product",
     [138, 139, 140], ["Pre & Post Workout Formula"]),
    ("firststring", "FirstString", "product",
     [141, 142, 143, 144], ["First String"]),
    ("wheysmooth", "WheySmooth", "product",
     [145, 146, 147, 148], ["WheySmooth"]),
    ("all-natural-wheysmooth", "All-Natural WheySmooth", "product",
     [149, 150, 151], ["WheySmooth"]),
    ("leanmeal", "LeanMeal", "product",
     [153, 154, 155, 156, 157, 160], ["LeanMeal Nutrition Shake"]),
    ("leanmeal-scripts", "LeanMeal — presentation and trainer script", "script",
     [158], ["LeanMeal Nutrition Shake"]),
    ("part-1-2-to-part-3", "Parts 1 & 2 leading to Part 3", "education",
     [162], []),
    ("wlls", "Weight Loss & Liver Support (WLLS)", "product",
     [163, 164, 165, 166, 167, 168], ["WeightLoss & LiverSupport"]),
    ("wlls-scripts", "WLLS — presentations and trainer/video script", "script",
     [169, 170], ["WeightLoss & LiverSupport"]),
    ("thermaccel", "ThermAccel", "product",
     [171, 172, 175, 176], ["ThermAccel"]),
    ("thermaccel-scripts", "ThermAccel — trainer scripts, pitch and flyer", "script",
     [173, 174, 177, 178, 179, 180], ["ThermAccel"]),
    ("carbrepel", "CarbRepel", "product",
     [181, 183, 184, 185], ["CarbRepel"]),
    ("carbrepel-scripts", "CarbRepel — presentations", "script",
     [182], ["CarbRepel"]),
    ("leanpak90", "Lean Pak 90 — the 90-day weight loss pack", "bundle",
     [186, 187, 188, 189, 190, 191, 192, 193, 194, 196, 197],
     ["LeanPak90", "WeightLoss & LiverSupport", "CarbRepel", "ThermAccel", "Alln1 SuperBlend", "Active MV",
      "LeanMeal Nutrition Shake"]),
    ("leanpak90-scripts", "Lean Pak 90 — presentation", "script",
     [195], ["LeanPak90"]),
    ("bodyfat-bundles", "Bodyfat / weight loss bundles, levels 1–3", "bundle",
     [198, 199, 200, 201, 202], ["LeanMeal Nutrition Shake", "Active MV", "Alln1 SuperBlend",
                                 "WeightLoss & LiverSupport", "ThermAccel", "CarbRepel", "AminoFormula",
                                 "Omega-3 Fish Oil"]),
    ("creatine-monohydrate", "Creatine Monohydrate", "product",
     [203, 204, 205, 206], ["Creatine Monohydrate"]),
    ("creatine-monohydrate-scripts", "Creatine Monohydrate — target-audience taglines", "script",
     [207, 208], ["Creatine Monohydrate"]),
    ("creatine-complex", "Creatine Complex (Extreme Creatine) with Careflow", "product",
     [209, 210, 211], ["Creatine Complex"]),
    ("creatine-complex-scripts", "Creatine Complex — taglines and flyer", "script",
     [212, 213], ["Creatine Complex"]),
    ("no7-preworkout", "NO7 PreWorkout", "product",
     [214, 215, 216, 218], ["NO7 PreWorkout"]),
    ("no7-preworkout-scripts", "NO7 PreWorkout — quick conversation", "script",
     [217], ["NO7 PreWorkout"]),
    ("aminoformula", "AminoFormula", "product",
     [219, 220, 221, 222, 223], ["AminoFormula"]),
    ("aminoformula-scripts", "AminoFormula — description, taglines and value proposition", "script",
     [224, 225], ["AminoFormula"]),
    ("glutamine-complex", "Glutamine Complex", "product",
     [227, 228], ["GlutamineComplex"]),
    ("muscle-gain-bundles", "Muscle / performance gain bundles, levels 1–3", "bundle",
     [229, 230, 232], ["Active MV", "Women's MV", "Alln1 SuperBlend", "Omega-3 Fish Oil", "Calcium Complex",
                       "WheySmooth", "First String", "Creatine Monohydrate", "AminoFormula",
                       "NO7 PreWorkout", "Creatine Complex"]),
    ("shred-bundles", "Shred bundles — Competitor and Champion", "bundle",
     [233, 234], ["Active MV", "Alln1 SuperBlend", "WeightLoss & LiverSupport", "LeanMeal Nutrition Shake",
                  "AminoFormula", "ThermAccel", "CarbRepel"]),
    ("ufc-challenge-bundles", "UFC Ultimate Challenge power bundle products", "bundle",
     [235], ["Active MV", "Women's MV", "Alln1 SuperBlend", "Creatine Monohydrate", "Creatine Complex",
             "Pre & Post Workout Formula", "WheySmooth", "AminoFormula", "NO7 PreWorkout"]),
    ("youth-nutrition", "Youth health and performance nutrition, levels 1 and 2", "bundle",
     [236, 238, 239, 240, 241, 242, 243, 244, 248, 249, 250, 252, 253, 254],
     ["Active MV", "Alln1 SuperBlend", "First String", "Plant Protein", "Calcium Complex",
      "Omega-3 Fish Oil", "Creatine Monohydrate"]),
    ("youth-nutrition-scripts", "Youth nutrition — promotion, trainer script sheet, website/email copy", "script",
     [245, 246, 247, 255], ["Active MV", "First String", "Creatine Monohydrate"]),
    ("family-wellness-bundle", "Family wellness bundle", "bundle",
     [15, 256, 257, 258, 259], ["Active MV", "Women's MV", "First String", "Omega-3 Fish Oil", "dotBAR",
                                "Vanilla dotWAFER", "Electrolytes"]),
    ("road-to-health-part-3", "The Road to Health — Part III: advanced therapies and GLP-1", "education",
     [260, 261, 262, 263, 264], ["LeanMeal Nutrition Shake", "Alln1 SuperBlend", "Active MV", "WheySmooth",
                                 "Plant Protein"]),
    ("road-to-health-part-3-career", "Part III — career enhancement summary", "script",
     [265], []),
    ("playspan-equals-lifespan", "Establishing a Playspan equal to a lifespan", "education",
     [266, 267, 268], []),
    ("playspan-equals-lifespan-script", "Playspan — trainer script", "script",
     [269], []),
]

# Slides in no section. Near-duplicates keep the most complete or newest
# wording (the deck carries several revisions of the same slide side by side).
DROPPED: dict[int, str] = {
    20: "older revision of 23",
    21: "subset of 23",
    22: "subset of 23",
    25: "duplicate of 24",
    29: "near-duplicate of 28",
    30: "older revision of 31",
    33: "duplicate of 30",
    34: "duplicate of 32",
    35: "duplicate of 1",
    38: "section divider (Health Products)",
    41: "near-duplicate of 40",
    57: "image of the text on 55-56",
    64: "near-duplicate of 61",
    65: "near-duplicate of 66 (per-day prices only)",
    71: "subset of 74",
    152: "competitor price comparison — prices only",
    159: "image of the text on 157",
    226: "title-only slide",
    231: "duplicate of 230",
    237: "older revision of 238",
    251: "duplicate of 248",
}

# What the bundle and image-only slides show, read off the rendered slides
# 2026-09-28. Product names as on the bottles; no prices.
PICTURED: dict[int, str] = {
    4: "First String; Creatine Monohydrate; NO7 PreWorkout; Omega-3 (as needed — add if <3 servings/wk "
       "fatty fish); Active MV OR Alln1 SuperBlend as the foundational dietary support.",
    5: "'All beginners and level 1 for muscle/performance gain': First String with Active MV OR Alln1 "
       "SuperBlend; Omega-3 as needed. 'Intermediate-advanced add-ons for muscle/performance gain': "
       "Creatine Monohydrate and/or NO7 PreWorkout. Creatine Monohydrate 'can also be included in any "
       "program seeking muscle preservation and cognitive support'.",
    11: "Two 'Family Health (and Performance) Cabinet' pictures. Foundation shelf: All Natural WheySmooth, "
        "Over50MV, Calcium Complex, Omega-3. 'Performance and Recovery' shelf: Creatine Monohydrate, "
        "AminoFormula, Collagen Complex. Family nutrition: dotBAR, dotWAFER. The smaller cabinet: "
        "First String, Active MV, Omega-3, Women's MV; Electrolytes; dotBAR, dotWAFER.",
    12: "Three senior bundles. Foundational: All Natural WheySmooth, Over50MV, Omega-3; Calcium Complex as "
        "needed to hit 1200 mg/d. Optimizing daily Performance: All Natural WheySmooth, Over50MV, Omega-3, "
        "Creatine Monohydrate. Beyond Protein to Recover & Build: All Natural WheySmooth, Over50MV, "
        "Omega-3, Creatine Monohydrate, AminoFormula.",
    13: "The same three senior bundles with Alln1 SuperBlend in place of Over50MV in each ('Substitute the "
        "Alln1 SuperBlend to amplify all outcomes'); Calcium Complex as needed to hit 1200 mg/d.",
    15: "'Our Family Health Cabinet'. Foundation: First String, Active MV, Omega-3, Women's MV. Daily health: "
        "Electrolytes. Family / convenient nutrition: dotBAR (Trail Mix, Strawberry, Fudge Graham), "
        "Vanilla dotWAFER.",
    16: "Beside each description, in order: Extreme Creatine (Creatine Complex); NO7 PreWorkout; Glutamine "
        "Complex; Digestive Enzymes; Sleep Aid; Electrolytes.",
    49: "An ActiveMV ad, text in the image: 'This is not a store-bought multivitamin. It's what lifelong "
        "exercisers use to stay active—for life. Training breaks you down. This makes rebuilding possible. "
        "You don't get stronger from workouts. You get stronger from recovery. And recovery only works if "
        "your body has the raw materials to rebuild. Most people train. Few supply the materials. That's "
        "why progress stalls. That's why joints ache. That's why \"aging\" feels sudden.'",
    50: "An ActiveMV ad, text in the image: 'This daily habit supports every workout you'll ever do. Why "
        "this formula is different? Built for active bodies; supports today's training & future aging; "
        "balanced—no megadoses, no gaps; works with food & other supplements; frequently updated as "
        "science evolves; this isn't generic, it's purpose-built. Why you won't find it in stores: retail "
        "multivitamins are made for mass appeal; practitioner-distributed; no watered-down formulas to cut "
        "costs; formulated by experts who understand the demands of an active body. The rule: take it "
        "daily. Keep taking it. This isn't a \"now\" supplement. It's a forever foundation. Your training "
        "tells your body where to adapt. This makes adaptation possible. Consistency matters more than "
        "intensity.'",
    85: "Multis 'for all life phases/genders': Active MV, Women's MV, Over50MV, OR Alln1 SuperBlend; "
        "Omega-3 (if needed); Calcium Complex (as needed to hit 1200 mg/d); All Natural WheySmooth; the "
        "protein line: First String, WheySmooth, Pre/Post Workout, LeanMeal, Plant Protein.",
    86: "Two Playspan essentials bundles: (1) All Natural WheySmooth, Omega-3, Active MV; (2) Alln1 "
        "SuperBlend, All Natural WheySmooth, Omega-3.",
    110: "The Digestive Enzymes label and bottle (90 capsules).",
    125: "Essentials: Active MV OR Alln1 SuperBlend, Omega-3, Vitamin D3 (as needed to hit >40 ng/ml). "
         "Next Level Performance Stabilizers & Enhancers: Brain Health, Creatine Monohydrate.",
    127: "'Essentials'. Protein: LeanMeal OR All Natural WheySmooth OR WheySmooth. Multi: Alln1 SuperBlend "
         "OR Active MV (Women's MV and Over50MV also shown).",
    128: "Five levels. 1 – Essentials ('3-Versions' of the multi): Active MV, Calcium Complex (as needed to "
         "hit 1200 mg/d), Omega-3; Vitamin D3 as needed to hit >30 ng/ml; Alln1 SuperBlend shown as the "
         "alternative. 2 – Living Stronger: adds Antioxidant. 3 – Stronger & Healthier: adds Probiotics. "
         "4 – … Staying Younger longer: adds Collagen Complex. 5 – … Thinking Better Longer: adds Brain "
         "Health. Recovery/longevity extras: Antioxidant, Probiotics, Collagen Complex, Brain Health; "
         "Digestive Enzymes when necessary to improve digestion & stomach comfort; Sleep Aid when necessary "
         "to improve sleep quality & duration.",
    129: "The SuperBlend version. Levels 1–3 (the Essentials through Stronger & Healthier): Alln1 SuperBlend, "
         "Calcium Complex (as needed to hit 1200 mg/d), Omega-3; Vitamin D3 as needed to hit >30 ng/ml. "
         "4: adds Collagen Complex. 5: adds Collagen Complex and Brain Health. Recovery/longevity extras: "
         "Antioxidant, Probiotics, Collagen Complex, Brain Health, Digestive Enzymes (when necessary), "
         "Sleep Aid (when necessary).",
    186: "The LeanPak90 kit (Weight Loss & Liver Support, ThermAccel, CarbRepel, weight loss planner, quick "
         "start card); the three bottles; images of GLP-1 drug pens.",
    187: "Weight Loss & Liver Support, ThermAccel, CarbRepel and the LeanPak90 kit ('All in 1 – Plus!').",
    196: "The LeanPak90 kit; Alln1 SuperBlend OR Active MV; the protein line (First String, WheySmooth, "
         "Pre/Post Workout, LeanMeal, Plant Protein).",
    198: "Three bodyfat reduction packages. 'Essentials to lose only bodyfat and protect muscle': LeanMeal, "
         "Active MV OR Alln1 SuperBlend. 'Accelerated bodyfat reduction with greater LBM protection & "
         "appetite control': LeanMeal, Active MV OR Alln1 SuperBlend, Weight Loss & Liver Support OR "
         "ThermAccel. 'Rapid bodyfat reduction while building muscle': LeanMeal, Active MV OR Alln1 "
         "SuperBlend, AminoFormula, Weight Loss & Liver Support OR ThermAccel. Omega-3 as needed.",
    199: "Collateral 'Weight Loss Bundles': Level 1 – LeanMeal, ActiveMV (or Alln1 SuperBlend), "
         "WeightLoss & LiverSupport. Level 2 – the same plus Amino Formula. Level 3 – the same plus "
         "CarbRepel or ThermAccel.",
    202: "A one-page summary as an image, '3 Levels of Body Fat Loss Bundles — Built for Your Goals, "
         "Experience & Budget': every bundle starts with the Active MVM and LeanMeal. Level 1 (Essentials, "
         "beginners or on a budget): MVM, LeanMeal, often adds WLLS for liver health with added appetite "
         "control. Level 2 (Accelerated, intermediate users or faster results): adds ThermAccel or "
         "CarbRepel. Level 3 (Comprehensive, advanced users, athletes, aggressive goals): includes "
         "AminoFormula. How to use: MVM daily with meals; LeanMeal replaces 1–2 meals a day within a "
         "structured 4–5 meal plan; Weight Loss & Liver Support 3 times daily; ThermAccel or CarbRepel as "
         "directed; AminoFormula pre-, during or post-workout.",
    230: "Three muscle/performance gain bundles. Level 1 – Build Fast (essentials & beginners): All Natural "
         "WheySmooth, Active MV OR Alln1 SuperBlend. Level-2 Build More Faster: WheySmooth, Active MV OR "
         "Alln1 SuperBlend, Creatine Monohydrate. Competitor Level (3): WheySmooth, Active MV OR Alln1 "
         "SuperBlend, Creatine Monohydrate, AminoFormula, &/OR Extreme Creatine (Creatine Complex). "
         "Omega-3 as needed.",
    232: "First String; Women's MV (18–50y, 1 tab) OR Active MV OR Alln1 SuperBlend; AminoFormula; "
         "Creatine Monohydrate &/or Extreme Creatine (Creatine Complex).",
    233: "The Champion Shred bundle: AminoFormula, LeanMeal, Active MV OR Alln1 SuperBlend, Weight Loss & "
         "Liver Support, CarbRepel OR ThermAccel.",
    235: "Active MV OR Women's MV OR Alln1 SuperBlend; Creatine Monohydrate and/or Extreme Creatine "
         "(Creatine Complex); Pre/Post Workout OR WheySmooth; AminoFormula (protein stacking); NO7 "
         "PreWorkout.",
    239: "Level-1 Essentials: First String, Active MV OR (≥12y) Alln1 SuperBlend, Omega-3 as needed, "
         "Calcium Complex (see calcium note). Level-2 (≥16yr with parent permission): First String, "
         "Active MV, Creatine Monohydrate.",
    240: "The ChildLife children's multivitamin (the linked dotFIT recommendation for children under 12); "
         "for 12–17y a dotFIT multi (1 tab) OR Alln1 SuperBlend; Calcium Complex; Omega-3; First String.",
    260: "GLP-1 drug images, plus the 'nutrition companion': LeanMeal, Alln1 SuperBlend OR Active MV, 'or "
         "favorite dF protein' (Plant Protein, WheySmooth — 'keeping it All-Natural or Vegan'); images of "
         "third-party anti-aging products (sermorelin, BPC-157, urolithin A and others).",
}


def repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def deck_path() -> Path:
    return repo_root() / "data" / "Product Summaries" / DECK_NAME


def out_dir() -> Path:
    return repo_root() / "processed" / "summaries"


def alias_path() -> Path:
    return repo_root() / "processed" / "aliases" / "alias_table.json"


def mask_prices(text: str) -> str:
    for pat in _PRICE_PATTERNS:
        text = pat.sub(PRICE, text)
    return text


def clean(text: str) -> str:
    """Vertical tabs (PowerPoint soft returns) to newlines, trailing spaces off."""
    text = text.replace("\x0b", "\n").replace("\xa0", " ")
    return "\n".join(ln.rstrip() for ln in text.split("\n")).strip()


def _pos(shape) -> tuple[int, int]:
    return (shape.top or 0, shape.left or 0)


def _run_link(run) -> str | None:
    try:
        addr = run.hyperlink.address
    except KeyError:  # internal slide jump: rId points at nothing external
        return None
    return addr or None


def read_slide(slide) -> dict:
    blocks: list[dict] = []
    links: list[dict] = []
    images: list[str] = []

    def walk(shapes) -> None:
        for s in sorted(shapes, key=_pos):
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(s.shapes)
                continue
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    images.append(hashlib.md5(s.image.blob).hexdigest()[:12])
                except (AttributeError, KeyError, ValueError):
                    pass  # linked, not embedded — nothing to hash
            if getattr(s, "has_table", False) and s.has_table:
                rows = [[mask_prices(clean(c.text)).replace("\n", " ")
                         for c in r.cells] for r in s.table.rows]
                if any(any(c for c in r) for r in rows):
                    blocks.append({"kind": "table", "rows": rows})
                continue
            if not s.has_text_frame:
                continue
            paras = []
            for para in s.text_frame.paragraphs:
                t = clean("".join(r.text for r in para.runs))
                if t:
                    paras.append({"level": para.level, "text": mask_prices(t)})
                for r in para.runs:
                    addr = _run_link(r)
                    if addr and r.text.strip():
                        links.append({"text": mask_prices(clean(r.text)),
                                      "url": addr})
            if paras:
                blocks.append({"kind": "text", "paras": paras})

    walk(slide.shapes)
    # one link per url: runs split mid-anchor repeat it ("In", "a", "90-", ...)
    seen: set[str] = set()
    uniq = []
    for ln in links:
        if ln["url"] not in seen:
            seen.add(ln["url"])
            uniq.append(ln)
    chars = sum(len(p["text"]) for b in blocks if b["kind"] == "text"
                for p in b["paras"])
    chars += sum(len(c) for b in blocks if b["kind"] == "table"
                 for r in b["rows"] for c in r)
    return {"blocks": blocks, "links": uniq, "images": images, "chars": chars}


def extract(path: Path) -> list[dict]:
    pres = Presentation(str(path))
    return [{"n": n, **read_slide(slide)} for n, slide in enumerate(pres.slides, 1)]


def block_md(block: dict) -> list[str]:
    if block["kind"] == "table":
        rows = [[c.replace("|", "\\|") for c in r] for r in block["rows"]]
        out = ["| " + " | ".join(rows[0]) + " |",
               "|" + "---|" * len(rows[0])]
        out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        return out
    out = []
    for p in block["paras"]:
        indent = "  " * p["level"]
        out.extend(indent + ln for ln in p["text"].split("\n"))
    return out


def slide_md(rec: dict, pictured: str | None = None) -> str:
    out = []
    for b in rec["blocks"]:
        out.extend(block_md(b))
        out.append("")
    if pictured:
        out.append(f"_Pictured (annotation, not slide text): {pictured}_")
        out.append("")
    if rec["links"]:
        out.append("Links: " + "; ".join(
            f"[{ln['text']}]({ln['url']})" for ln in rec["links"]))
        out.append("")
    return "\n".join(out).rstrip()


def validate(n_slides: int, families: dict[str, list[int]]) -> list[str]:
    """Every slide in exactly one section or dropped; every product resolves."""
    errors: list[str] = []
    placed: dict[int, str] = {}
    slugs: set[str] = set()
    for slug, _title, kind, slides, products in SECTIONS:
        if slug in slugs:
            errors.append(f"duplicate section slug {slug}")
        slugs.add(slug)
        if kind not in KINDS:
            errors.append(f"{slug}: unknown kind {kind!r}")
        for n in slides:
            if n in placed:
                errors.append(f"slide {n} in both {placed[n]} and {slug}")
            placed[n] = slug
        errors.extend(f"{slug}: product {p!r} is not an alias-table family"
                      for p in products if p not in families)
    for n in DROPPED:
        if n in placed:
            errors.append(f"slide {n} dropped but also in {placed[n]}")
    missing = [n for n in range(1, n_slides + 1) if n not in placed and n not in DROPPED]
    if missing:
        errors.append(f"slides in no section and not dropped: {missing}")
    stray = [n for n in list(placed) + list(DROPPED) + list(PICTURED) if not 1 <= n <= n_slides]
    if stray:
        errors.append(f"slide numbers outside the deck: {sorted(set(stray))}")
    return errors


def render_section(idx: int, section, by_n: dict[int, dict],
                   families: dict[str, list[int]]) -> str:
    slug, title, kind, slides, products = section
    tagged = ", ".join(f"{p} [{'/'.join(map(str, families[p]))}]" for p in products)
    out = [f"# {title}", "",
           f"- source: `{DECK_NAME}`, slides {', '.join(map(str, slides))}",
           f"- kind: {kind}",
           f"- products: {tagged or '—'}", ""]
    for n in slides:
        out.append(f"## Slide {n}")
        out.append("")
        body = slide_md(by_n[n], PICTURED.get(n))
        out.append(body if body else "_(no text)_")
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def render_dump(slides: list[dict]) -> str:
    out = [f"# {DECK_NAME} — all slides, text-exact, prices masked", ""]
    for rec in slides:
        out.append(f"## Slide {rec['n']} ({rec['chars']} chars, {len(rec['images'])} images)")
        out.append("")
        out.append(slide_md(rec, PICTURED.get(rec["n"])) or "_(no text)_")
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    path = deck_path()
    if not path.exists():
        print(f"FAIL: deck not found: {path}", file=sys.stderr)
        return 2
    table = json.loads(alias_path().read_text(encoding="utf-8"))
    families = {f["family"]: f["part_nos"] for f in table["families"]}

    slides = extract(path)
    errors = validate(len(slides), families)
    if errors:
        for e in errors:
            print(f"FAIL: {e}", file=sys.stderr)
        return 1

    od = out_dir()
    md_dir = od / "md"
    md_dir.mkdir(parents=True, exist_ok=True)
    for stale in md_dir.glob("*.md"):
        stale.unlink()
    by_n = {s["n"]: s for s in slides}
    meta = []
    for i, section in enumerate(SECTIONS, 1):
        slug, title, kind, sl, products = section
        name = f"{i:02d}-{slug}.md"
        (md_dir / name).write_text(render_section(i, section, by_n, families), encoding="utf-8")
        meta.append({"file": name, "slug": slug, "title": title, "kind": kind,
                     "slides": sl, "products": products,
                     "part_nos": sorted({p for f in products for p in families[f]}),
                     "chars": sum(by_n[n]["chars"] for n in sl),
                     "pictured": [n for n in sl if n in PICTURED]})
    masked = sum(slide_md(s).count(PRICE) for s in slides)
    summary = {
        "deck": DECK_NAME, "n_slides": len(slides), "n_sections": len(meta),
        "by_kind": {k: sum(1 for m in meta if m["kind"] == k) for k in KINDS},
        "prices_masked": masked,
        "dropped": {str(n): r for n, r in sorted(DROPPED.items())},
        "sections": meta,
    }
    (od / "summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    runs = od / "runs"
    runs.mkdir(exist_ok=True)
    (runs / "slides.md").write_text(render_dump(slides), encoding="utf-8")

    print(f"done: {len(slides)} slides -> {len(meta)} sections "
          f"({', '.join(f'{k} {v}' for k, v in summary['by_kind'].items())}), "
          f"{len(DROPPED)} dropped, {masked} prices masked -> {md_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
